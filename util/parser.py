from pathlib import Path
import pypdf
import docx

def text_extractor(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()
    extracted_text = ""

    if extension == '.pdf':
        reader = pypdf.PdfReader(file_path)
        for page in reader.pages:
            extracted_text += (page.extract_text() or "") + '\n' or ""

    elif extension == '.docx':
        doc = docx.Document(file_path)
        for paragraph in doc.paragraphs:
            extracted_text += paragraph.text + '\n' or ""
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    extracted_text += ' | '.join(row_text) + '\n'

    elif extension == '.txt':
        with open(file_path, 'r', encoding = 'utf-8') as f:
            extracted_text = f.read()

    else:
        raise ValueError (f'Unsupported file format: {extension}')
    return extracted_text.strip()


if __name__ == '__main__':
    sample_path = 'test_sample.txt'
    with open(sample_path, 'w', encoding = 'utf-8') as f:
        f.write('Position: AI Engineer\nRequirements: Python, FastAPI, Groq API')

    try:
        extracted = text_extractor(sample_path)
        print('---- Extracted Output ----')
        print(extracted)
        print('--------')
        print('SUCCESS: Text Extractor is fucking bomb!')

    except Exception as e:
        print(f'ERROR: {e}')